"""sb-lowfi-pptx: low-fi 와이어프레임 geo.json -> 실무형 화면설계서 PPTX

구성
- 표지 -> 화면 목록 -> [PC 화면] 구분 슬라이드 -> PC 시나리오 전체 -> [Mobile 화면] 구분 슬라이드 -> Mobile 시나리오 전체
- 화면 슬라이드: 상단 메타 표 / 좌측 와이어프레임(플랫폼별 고정 축척) / 우측 Description(이 슬라이드의 마커 번호만)
- 상태별 화면·처리 규칙: 화면 슬라이드에 함께 들어가면 함께, 넘치면 바로 다음 슬라이드
- 고정 축척으로 한 장에 들어가지 않는 긴 화면은 글자가 지나가지 않는 빈 줄에서 상단/하단으로 나눔
- 마커는 글자·다른 마커와 겹치지 않는 자리로 자동 배치

사용: build_pptx.py --geo geo.json --out 화면설계서_YYYYMMDD_HHMM.pptx [--md 화면설계서.md] [--project 이름] [--date] [--version] [--note]
"""
import json, re, math, argparse, datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

ap = argparse.ArgumentParser()
ap.add_argument('--geo', required=True); ap.add_argument('--out', required=True)
ap.add_argument('--md', default=None)
ap.add_argument('--project', default='Project'); ap.add_argument('--date', default=datetime.date.today().isoformat())
ap.add_argument('--version', default='v1.0'); ap.add_argument('--note', default='')
ap.add_argument('--pc-scale', type=float, default=0.0062, help='PC 프레임 1px당 인치 (1000px -> 6.2in)')
ap.add_argument('--mo-scale', type=float, default=0.0068, help='Mobile 프레임 1px당 인치 (375px -> 2.55in)')
A = ap.parse_args()
OUT, DATE, VER, PROJ = A.out, A.date, A.version, A.project
geo = json.load(open(A.geo, encoding='utf-8'))
md = open(A.md, encoding='utf-8').read() if A.md else ''

FONT = '맑은 고딕'
INK = RGBColor(0x22, 0x22, 0x22); GRAY = RGBColor(0x6b, 0x6b, 0x73); ORANGE = RGBColor(0xe8, 0x59, 0x0c)
HDR = RGBColor(0x40, 0x40, 0x40); LIGHT = RGBColor(0xf4, 0xf4, 0xf5); LINE = RGBColor(0x9a, 0x9a, 0xa2); WHITE = RGBColor(255, 255, 255)
SW, SH = 13.333, 7.5
L, R = 0.15, SW - 0.15          # 좌우 여백
TOP = 1.55                      # 메타 표 아래 콘텐츠 시작 (고정)
BOT = 7.14                      # 콘텐츠 끝 (바닥글 위)
CAP = 0.24                      # 프레임 캡션 높이
DESC_SIZE = 8                   # Description 글자 크기 (고정)
MK = 0.19                       # 마커 지름 (in)
K = {'pc': A.pc_scale, 'mo': A.mo_scale}
FRAME_H = BOT - TOP - CAP       # 프레임이 쓸 수 있는 높이
LAYOUT = {  # 플랫폼별 고정 레이아웃: 와이어프레임 영역 폭, Description 시작 x
    'pc': {'area_w': 1000 * K['pc'], 'desc_x': L + 1000 * K['pc'] + 0.3},
    'mo': {'area_w': 375 * K['mo'] * 3 + 0.2 * 2, 'desc_x': L + 375 * K['mo'] * 3 + 0.2 * 2 + 0.3},
}
prs = Presentation(); prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
blank = prs.slide_layouts[6]

# ---------- 원본 md 메타 ----------
meta = {}
blocks_md = re.split(r'\n(?=### SCR-)', md) if md else []
for blk in blocks_md:
    m = re.match(r'### (SCR-\d{3})\. (.+)', blk)
    if not m: continue
    d = {'name': m.group(2).strip()}
    for k in ['URL', '플랫폼', '접근 권한', '관련 기능', '관련 API']:
        r = re.search(r'^\| %s \| (.+?) \|\s*$' % re.escape(k), blk, re.M)
        d[k] = r.group(1).strip() if r else '-'
    meta[m.group(1)] = d
for scr in geo:
    S = scr['sid'].upper()
    if S not in meta:
        parts = [p.strip() for p in scr.get('meta', '').split('·')]
        meta[S] = {'name': re.sub(r'^SCR-\d+\s*', '', scr['h2']), 'URL': parts[0] if parts else '-', '플랫폼': parts[1] if len(parts) > 1 else '-',
                   '접근 권한': parts[2] if len(parts) > 2 else '-', '관련 기능': parts[3] if len(parts) > 3 else '-', '관련 API': ' · '.join(parts[4:]) or '-'}
prio = {m.group(1): m.group(2) for m in re.finditer(r'^\| (SCR-\d{3}) \|.*?\| (P\d) \|', md, re.M)}

def md_stbls(sid):
    blk = next((b for b in blocks_md if b.startswith('### %s.' % sid)), '')
    out = []
    m = re.search(r'\*\*상태별 화면\*\*\s*\n((?:\|.*\n?)+)', blk)
    if m:
        rows = [[c.strip() for c in l.strip().strip('|').split('|')] for l in m.group(1).strip().split('\n') if not re.match(r'^\|[-\s|]+\|$', l.strip())]
        out.append({'title': '상태별 화면', 'rows': rows, 'items': []})
    r = re.search(r'\*\*처리 규칙\*\*\s*\n((?:\s*[-*] .*\n?)+)', blk)
    if r:
        out.append({'title': '처리 규칙', 'rows': [], 'items': [re.sub(r'^\s*[-*] ', '', x).strip() for x in r.group(1).strip().split('\n') if x.strip()]})
    return out

for scr in geo:
    S = scr['sid'].upper()
    have = {('r' if st['rows'] else 'i') for st in scr['stbls']}
    for st in md_stbls(S):
        if ('r' if st['rows'] else 'i') not in have: scr['stbls'].append(st)

# ---------- 공통 유틸 ----------
def rgb(c):
    m = re.findall(r'[\d.]+', c); r, g, b = [float(x) for x in m[:3]]; a = float(m[3]) if len(m) > 3 else 1
    if a < 1: r, g, b = [255 * (1 - a) + v * a for v in (r, g, b)]
    return RGBColor(int(r), int(g), int(b))

def wch(t): return sum(1.0 if ord(c) > 0x2E80 else 0.56 for c in t)

def tb(sl, x, y, w, h, text, size, color=INK, bold=False, wrap=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, font=FONT):
    s = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = s.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = font; f.color.rgb = color
    return s

def rect(sl, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); s.shadow.inherit = False
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(lw)
    return s

def set_border(ce, color='808080', w=6350):
    tcPr = ce._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for old in tcPr.findall(qn(tag)): tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag), w=str(w), cap='flat', cmpd='sng', algn='ctr')
        sf = etree.SubElement(ln, qn('a:solidFill')); etree.SubElement(sf, qn('a:srgbClr'), val=color)
        etree.SubElement(ln, qn('a:prstDash'), val='solid')
    for f in [c for c in tcPr if c.tag in (qn('a:solidFill'), qn('a:noFill'))]: tcPr.remove(f); tcPr.append(f)

def cell_set(ce, text, size, bold=False, color=INK, fill=None, align=PP_ALIGN.LEFT):
    ce.margin_left = ce.margin_right = Inches(0.06); ce.margin_top = ce.margin_bottom = Inches(0.025)
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = ce.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = color
    ce.fill.solid(); ce.fill.fore_color.rgb = fill or WHITE
    set_border(ce)

def new_table(sl, rows, cols, x, y, w, h):
    gt = sl.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tblPr = gt._element.graphic.graphicData.tbl.tblPr; tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
    return gt.table

# ---------- 와이어프레임 그리기 (마커 겹침 회피 포함) ----------
def text_geom(it, k, ox, oy):
    single = (it['lines'] <= 1) or (it['h'] < it['size'] * 1.9)
    pt = max(4.5, it['size'] * k * 72)
    tw = max(it['w'] * k * 1.04, wch(it['text']) * pt / 72 * 0.98 if single else 0)   # 실제 글자 폭 추정 (왼쪽 시작점은 정확, 오른쪽은 글꼴 차이 여유)
    bw = max(it['w'] * k * (1.18 if single else 1.06), wch(it['text']) * pt / 72 * 1.12 if single else 0)  # 텍스트 상자 폭
    x = it['x'] * k
    al = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(it['align'], PP_ALIGN.LEFT)
    if al == PP_ALIGN.CENTER: bx, tx = x - (bw - it['w'] * k) / 2, x - (tw - it['w'] * k) / 2
    elif al == PP_ALIGN.RIGHT: bx, tx = x - (bw - it['w'] * k), x - (tw - it['w'] * k)
    else: bx, tx = x, x
    return dict(single=single, pt=pt, al=al, box=(ox + bx, oy + it['y'] * k, bw, it['h'] * k), ink=(ox + tx, oy + it['y'] * k, tw, it['h'] * k))

def overlap(a, b, pad=0.0):
    return not (a[0] + a[2] <= b[0] - pad or b[0] + b[2] <= a[0] - pad or a[1] + a[3] <= b[1] - pad or b[1] + b[3] <= a[1] - pad)

def place_markers(marks, texts, k, ox, oy, fw, fh, y0):
    """marks: 마커 item, texts: 글자 영역(in). 반환: [(n, cx, cy)] 겹치지 않는 중심 좌표"""
    placed = []; d = MK; g = 0.03
    lo_x, hi_x, lo_y, hi_y = ox - 0.12, ox + fw * k + 0.12, oy - 0.12, oy + fh * k + 0.05
    for m in sorted(marks, key=lambda m: (m['y'], m['x'])):
        ex, ey, ew, eh = ox + m.get('ex', m['x'] + 8) * k, oy + (m.get('ey', m['y'] + 8) - y0) * k, m.get('ew', 20) * k, m.get('eh', 20) * k
        orig = (ox + (m['x'] + m['w'] / 2) * k, oy + (m['y'] - y0 + m['h'] / 2) * k)
        cands = [orig,
                 (ex - d / 2 - g, ey + min(eh, d) / 2),            # 요소 왼쪽 바깥, 위쪽 정렬
                 (ex - d / 2 - g, ey + eh / 2),                     # 요소 왼쪽 바깥, 가운데
                 (ex + d / 2, ey - d / 2 - g),                      # 요소 위 바깥, 왼쪽
                 (ex + ew + d / 2 + g, ey + min(eh, d) / 2),       # 요소 오른쪽 바깥
                 (ex + ew - d / 2, ey - d / 2 - g),                 # 요소 위 바깥, 오른쪽
                 (ex + d / 2 + g, ey + d / 2 + g),                  # 요소 안쪽 왼쪽 위
                 (ex + ew - d / 2 - g, ey + d / 2 + g)]             # 요소 안쪽 오른쪽 위
        for step in range(1, 7):                                    # 그래도 없으면 바깥으로 조금씩
            cands.append((ex - d / 2 - g - step * (d + g), ey + min(eh, d) / 2))
            cands.append((ex + d / 2 + step * (d + g), ey - d / 2 - g))
        n_step = max(1, int(ew / (d / 2)))                          # 마지막: 요소 안쪽·위·아래 가장자리를 훑어 빈자리 찾기
        for i in range(n_step + 1):
            cx = ex + d / 2 + i * (ew - d) / n_step
            for cy in (ey + eh / 2, ey + d / 2 + g, ey + eh - d / 2 - g, ey - d / 2 - g, ey + eh + d / 2 + g):
                cands.append((cx, cy))
        best = None
        for cx, cy in cands:
            r = (cx - d / 2, cy - d / 2, d, d)
            if r[0] < lo_x or r[0] + d > hi_x or r[1] < lo_y or r[1] + d > hi_y: continue
            if any(overlap(r, t, 0.008) for t in texts): continue
            if any(overlap(r, (px - d / 2, py - d / 2, d, d), g) for _, px, py in placed): continue
            best = (cx, cy); break
        if best is None: best = orig; COLLIDE.append((CUR[0], m['n']))
        placed.append((m['n'], best[0], best[1]))
    return placed

COLLIDE = []; CUR = ['']
def draw(sl, fr, ox, oy, k, y0=0, y1=None):
    """프레임 fr을 (ox, oy)에 축척 k로 그림. y0~y1(px) 구간만 (긴 화면 분할용). 반환: 표시한 마커 번호"""
    if y1 is None: y1 = fr['h']
    CUR[0] = fr['cap'][:30]
    X = lambda v: ox + v * k
    texts = []; marks = []
    for it in fr['items']:
        t = it['t']
        top, bottom = it['y'], it['y'] + it['h']
        if t == 'mark':
            c = it.get('ey', it['y']) + min(it.get('eh', 20), 30) / 2   # 요소 윗부분이 보이는 구간에 마커를 둠
            if y0 <= c < y1: marks.append(it)
            continue
        if bottom <= y0 or top >= y1: continue
        if t == 'box':
            yy, hh = max(top, y0), min(bottom, y1) - max(top, y0)
            w, h = it['w'] * k, hh * k
            if w < 0.005 or h < 0.005: continue
            bw = it['bw']; allsame = len(set(bw)) == 1 and bw[0] > 0
            rad = it['rad']; mn = min(it['w'], it['h'])
            if it['pct']: rad = mn * rad / 100
            shape = MSO_SHAPE.RECTANGLE
            if rad >= mn / 2 - 0.5 and rad > 0 and abs(it['w'] - it['h']) < 2: shape = MSO_SHAPE.OVAL
            elif rad > 1: shape = MSO_SHAPE.ROUNDED_RECTANGLE
            s = sl.shapes.add_shape(shape, Inches(X(it['x'])), Inches(oy + (yy - y0) * k), Inches(w), Inches(h)); s.shadow.inherit = False
            if shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = min(0.5, rad / max(1, min(it['w'], hh)))
            if it['bg']: s.fill.solid(); s.fill.fore_color.rgb = rgb(it['bg'])
            else: s.fill.background()
            if allsame:
                s.line.color.rgb = rgb(it['bc'][0]); s.line.width = Pt(max(0.5, bw[0] * k * 72))
                if it['bs'] == 'dashed': s.line.dash_style = MSO_LINE.DASH
            else:
                s.line.fill.background()
                for i, b in enumerate(bw):
                    if b <= 0: continue
                    if i == 0 and top < y0: continue
                    if i == 2 and bottom > y1: continue
                    x0, ya, x1, yb = X(it['x']), oy + (yy - y0) * k, X(it['x'] + it['w']), oy + (yy - y0 + hh) * k
                    seg = [(x0, ya, x1, ya), (x1, ya, x1, yb), (x0, yb, x1, yb), (x0, ya, x0, yb)][i]
                    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *[Inches(v) for v in seg])
                    c.line.color.rgb = rgb(it['bc'][i]); c.line.width = Pt(max(0.5, b * k * 72))
        elif t == 'text':
            if top < y0 or bottom > y1: continue      # 분할선은 글자가 없는 곳에서만 자르므로 걸치는 글자는 없음
            gm = text_geom(it, k, ox, oy - y0 * k)
            bx, by, bw_, bh = gm['box']
            s = tb(sl, bx, by, bw_, bh, it['text'], gm['pt'], rgb(it['color']), it['bold'], wrap=not gm['single'], align=gm['al'], italic=it['italic'], font=('Consolas' if it['mono'] else FONT))
            if it['ul']: s.text_frame.paragraphs[0].runs[0].font.underline = True
            texts.append(gm['ink'])
    shown = []
    for n, cx, cy in place_markers(marks, texts, k, ox, oy, fr['w'], y1 - y0, y0):
        s = rect(sl, cx - MK / 2, cy - MK / 2, MK, MK, fill=ORANGE, shape=MSO_SHAPE.OVAL)
        tf = s.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(n)
        r.font.size = Pt(7 if len(str(n)) < 2 else 6.5); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
        shown.append(str(n))
    return shown

def cut_points(fr, maxh):
    """고정 축척에서 maxh(px)를 넘는 프레임의 분할 y 목록. 글자가 걸치지 않는 y 중 가장 아래를 고른다"""
    busy = [(it['y'], it['y'] + it['h']) for it in fr['items'] if it['t'] == 'text']
    def free(y): return all(not (a - 2 < y < b + 2) for a, b in busy)
    cuts, start = [], 0
    while fr['h'] - start > maxh:
        y = int(start + maxh)
        while y > start + maxh * 0.5 and not free(y): y -= 2
        if y <= start + maxh * 0.5: y = int(start + maxh)
        cuts.append(y); start = y
    return cuts

# ---------- 메타 표 ----------
PAGE = [0]
HCOLW = [1.0, 4.85, 0.95, 3.05, 0.8, 2.05]
def header_rows(S, kind, part):
    m = meta.get(S, {})
    view = 'PC 화면' if kind == 'pc' else 'Mobile 화면'
    return [
        [('Page Title', 1), (m.get('name', '') + (f'  ({part})' if part else ''), 1), ('Screen ID', 1), (S, 1), ('Date', 1), (DATE, 1)],
        [('Screen Path', 1), (m.get('URL', '-'), 1), ('Platform', 1), (f"{m.get('플랫폼', '-')} / {prio.get(S, 'P0')}  (이 슬라이드: {view})", 3)],
        [('접근 권한', 1), (m.get('접근 권한', '-'), 5)],
        [('관련 기능', 1), (m.get('관련 기능', '-'), 5)],
        [('관련 API', 1), (m.get('관련 API', '-'), 5)],
    ]

def row_lines(row, size):
    c, n = 0, 1
    for j, (txt, span) in enumerate(row):
        w = sum(HCOLW[c:c + span]) - 0.12
        if j % 2: n = max(n, math.ceil(wch(txt) / max(8, w * 72 / (size * 0.95))))
        c += span
    return n

def slide_base(S, kind, part=None):
    sl = prs.slides.add_slide(blank); PAGE[0] += 1
    rows = header_rows(S, kind, part)
    size = 7.5
    while size > 6 and 0.1 + sum(0.245 + 0.125 * (row_lines(r, size) - 1) for r in rows) > TOP - 0.08: size -= 0.5
    tbl = new_table(sl, len(rows), 6, L, 0.1, sum(HCOLW), 0.2 * len(rows))
    for i, w in enumerate(HCOLW): tbl.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(0.245 + 0.125 * (row_lines(row, size) - 1))
        c = 0
        for j, (txt, span) in enumerate(row):
            ce = tbl.cell(r, c)
            if span > 1: ce.merge(tbl.cell(r, c + span - 1))
            lab = j % 2 == 0
            cell_set(ce, txt, size, bold=lab, color=WHITE if lab else INK, fill=HDR if lab else None, align=PP_ALIGN.CENTER if lab else PP_ALIGN.LEFT)
            c += span
    tb(sl, L, 7.24, 6, 0.2, f'{PROJ} 화면설계서 {VER}  ·  low-fidelity wireframe', 7, GRAY)
    tb(sl, R - 1.0, 7.24, 1.0, 0.2, f'{PAGE[0]} page', 7, GRAY, align=PP_ALIGN.RIGHT)
    return sl

# ---------- Description ----------
def rng(s):
    nums = set()
    for a in re.findall(r'\d+(?:\s*[~-]\s*\d+)?', s):
        p = re.split(r'\s*[~-]\s*', a)
        nums |= set(range(int(p[0]), int(p[1]) + 1)) if len(p) == 2 else {int(p[0])}
    return nums

def comp_blocks(scr, shown):
    keep = set(int(x) for x in shown)
    rows = [r for r in scr['spec'] if rng(r[0]) & keep]
    out = [('h', '', '화면 구성 요소')]
    for r in rows:
        out.append(('c', r[0], re.sub(r'\s+', ' ', r[1]) + ': ' + re.sub(r'\s+', ' ', r[2] if len(r) > 2 else '')))
    return out, len(rows)

def state_blocks(scr):
    out = []
    for st in scr['stbls']:
        if st['rows']:
            out.append(('h', '', st['title'] or '상태별 화면'))
            for row in st['rows'][1:]:
                out.append(('s', '', f'{row[0]}  |  {row[1]}  ->  {row[2]}' if len(row) >= 3 else ' | '.join(row)))
        if st['items']:
            out.append(('h', '', st['title'] or '처리 규칙'))
            for it in st['items']: out.append(('r', '', it))
    return out

def blk_h(b, width, size):
    cpl = max(8, (width - (0.3 if b[0] == 'c' else 0.12)) * 72 / (size * 0.92))
    lines = max(1, math.ceil(wch(b[2]) / cpl))
    return lines * size * 1.32 / 72 + (0.09 if b[0] == 'h' else 0.035)

def total_h(blocks, width, size): return sum(blk_h(b, width, size) for b in blocks)

def split_cols(blocks, width, height, size):
    cols, cur, h = [], [], 0
    for b in blocks:
        bh = blk_h(b, width, size)
        if cur and (h + bh > height or (b[0] == 'h' and h + bh + blk_h(('s', '', 'x' * 40), width, size) > height)):
            cols.append(cur); cur, h = [], 0
        cur.append(b); h += bh
    if cur: cols.append(cur)
    return cols

def draw_blocks(sl, ix, cy, iw, blocks, size):
    for b in blocks:
        bh = blk_h(b, iw, size)
        if b[0] == 'h':
            rect(sl, ix, cy + 0.01, iw, size * 1.32 / 72 + 0.04, fill=LIGHT)
            tb(sl, ix + 0.05, cy + 0.03, iw - 0.1, size * 1.32 / 72, b[2], size + 0.5, INK, True, wrap=False)
        elif b[0] == 'c':
            no = b[1]
            if len(no) <= 2:
                o = rect(sl, ix, cy + 0.005, 0.17, 0.17, fill=ORANGE, shape=MSO_SHAPE.OVAL)
                otf = o.text_frame; otf.margin_left = otf.margin_right = otf.margin_top = otf.margin_bottom = 0; otf.word_wrap = False; otf.vertical_anchor = MSO_ANCHOR.MIDDLE
                op = otf.paragraphs[0]; op.alignment = PP_ALIGN.CENTER; orr = op.add_run(); orr.text = no
                orr.font.size = Pt(6.5); orr.font.bold = True; orr.font.color.rgb = WHITE; orr.font.name = FONT
            else:
                tb(sl, ix, cy, 0.3, 0.17, no, size - 0.5, ORANGE, True, wrap=False)
            s = tb(sl, ix + 0.3, cy, iw - 0.3, bh, '', size)
            p = s.text_frame.paragraphs[0]; nm, _, rest = b[2].partition(': ')
            r1 = p.runs[0]; r1.text = nm; r1.font.bold = True
            r2 = p.add_run(); r2.text = '  ' + rest; r2.font.size = Pt(size); r2.font.name = FONT; r2.font.color.rgb = INK
        elif b[0] == 'n':
            tb(sl, ix + 0.05, cy, iw - 0.05, bh, b[2], size, GRAY)
        else:
            s = tb(sl, ix + 0.05, cy, iw - 0.05, bh, '', size, INK); p = s.text_frame.paragraphs[0]
            if b[0] == 's':
                st, _, rest = b[2].partition('  |  ')
                p.runs[0].text = '· '; r1 = p.add_run(); r1.text = st; r1.font.bold = True; r1.font.size = Pt(size); r1.font.name = FONT; r1.font.color.rgb = INK
                r2 = p.add_run(); r2.text = '  |  ' + rest; r2.font.size = Pt(size); r2.font.name = FONT; r2.font.color.rgb = GRAY
            else:
                p.runs[0].text = '- ' + b[2]
        cy += bh

def panel(sl, x, w, title, blocks_cols, size):
    y = TOP + 0.03
    hb = rect(sl, x, y, w, 0.26, fill=HDR)
    tf = hb.text_frame; tf.margin_top = tf.margin_bottom = 0; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    rect(sl, x, y + 0.26, w, BOT - y - 0.26, line=LINE)
    n = max(1, len(blocks_cols)); cw = (w - 0.2 - 0.25 * (n - 1)) / n
    for i, col in enumerate(blocks_cols): draw_blocks(sl, x + 0.1 + i * (cw + 0.25), y + 0.36, cw, col, size)

PANEL_H = BOT - TOP - 0.45

# ---------- 화면 슬라이드 ----------
def screen_slides(scr, kind):
    """이 화면의 kind 프레임들을 고정 축척으로 슬라이드에 나눠 담는다. 반환: [[(frame, y0, y1, caption)]]"""
    k = K[kind]; maxh = FRAME_H / k
    units = []
    for f in [f for f in scr['frames'] if f['kind'] == kind]:
        cuts = cut_points(f, maxh)
        if not cuts: units.append((f, 0, f['h'], f['cap'])); continue
        ys = [0] + cuts + [f['h']]
        for i in range(len(ys) - 1):
            units.append((f, ys[i], ys[i + 1], f"{f['cap']} ({'상단' if i == 0 else ('하단' if i == len(ys) - 2 else '중단')})"))
    groups, cur = [], []
    for u in units:
        if kind == 'pc':
            used = sum((x[2] - x[1]) * k + CAP for x in cur)
            fits = used + (u[2] - u[1]) * k + CAP <= BOT - TOP + 0.001
        else:
            fits = len(cur) < 3
        if cur and not fits: groups.append(cur); cur = []
        cur.append(u)
    if cur: groups.append(cur)
    return groups

def render_screen(scr, kind, pages_of):
    S = scr['sid'].upper(); k = K[kind]; lay = LAYOUT[kind]
    groups = screen_slides(scr, kind)
    if not groups: return
    dx = lay['desc_x']; dw = R - dx
    states = state_blocks(scr)
    pending = []   # 다음 슬라이드로 넘길 블록
    for gi, grp in enumerate(groups):
        part = ('PC' if kind == 'pc' else 'Mobile') + (f' {gi + 1}/{len(groups)}' if len(groups) > 1 else '')
        sl = slide_base(S, kind, part)
        pages_of.setdefault((S, kind), PAGE[0])
        shown = []
        if kind == 'pc':
            y = TOP
            for f, y0, y1, cap in grp:
                tb(sl, L, y + 0.02, lay['area_w'], 0.2, cap or 'PC', 8, GRAY, True, wrap=False)
                shown += draw(sl, f, L, y + CAP, k, y0, y1); y += (y1 - y0) * k + CAP
        else:
            x = L
            for f, y0, y1, cap in grp:
                tb(sl, x, TOP + 0.02, 375 * k, 0.2, cap or 'Mobile', 8, GRAY, True, wrap=False)
                shown += draw(sl, f, x, TOP + CAP, k, y0, y1); x += 375 * k + 0.2
        comp, ncomp = comp_blocks(scr, shown)
        if not ncomp:
            caps = ', '.join(dict.fromkeys(c for _, _, _, c in grp))
            comp = [('h', '', '화면 구성 요소'), ('n', '', f'이 슬라이드는 상태 화면입니다 ({caps}). 번호별 설명은 같은 화면의 기본 슬라이드를 참고하세요.')]
        last = gi == len(groups) - 1
        body = comp + (states if last else [])
        if total_h(body, dw - 0.2, DESC_SIZE) <= PANEL_H:
            panel(sl, dx, dw, 'Description', [body], DESC_SIZE)
        else:
            cols = split_cols(comp, dw - 0.2, PANEL_H, DESC_SIZE)
            panel(sl, dx, dw, 'Description', cols[:1], DESC_SIZE)
            rest = [b for c in cols[1:] for b in c]
            if rest: pending += [('h', '', f'화면 구성 요소 (계속, {part})')] + rest
            if last: pending += states
    if pending: state_slide(scr, kind, pending)

def state_slide(scr, kind, blocks):
    S = scr['sid'].upper()
    width = R - L
    cw = (width - 0.2 - 0.25) / 2
    cols = split_cols(blocks, cw, PANEL_H, DESC_SIZE)
    for i in range(0, len(cols), 2):
        sl = slide_base(S, kind, ('PC' if kind == 'pc' else 'Mobile') + ' · 상태별 화면·처리 규칙' + (f' {i // 2 + 1}' if len(cols) > 2 else ''))
        panel(sl, L, width, 'Description: 상태별 화면 · 처리 규칙', cols[i:i + 2], DESC_SIZE)

# ---------- 표지·목차·구분 ----------
def cover():
    sl = prs.slides.add_slide(blank); PAGE[0] += 1
    rect(sl, 0, 3.1, SW, 0.5, fill=HDR)
    tb(sl, 0, 2.2, SW, 0.7, PROJ, 36, INK, True, align=PP_ALIGN.CENTER)
    tb(sl, 0, 3.17, SW, 0.4, '화면설계서 (Low-fidelity Wireframe)', 16, WHITE, True, align=PP_ALIGN.CENTER)
    tb(sl, 0, 4.0, SW, 0.3, f'{VER}  ·  {DATE}  ·  화면 {len(geo)}개 (PC 시나리오 → Mobile 시나리오)', 11, GRAY, align=PP_ALIGN.CENTER)
    if A.note: tb(sl, 0, 4.35, SW, 0.3, A.note, 9, GRAY, align=PP_ALIGN.CENTER)

def divider(title, sub):
    sl = prs.slides.add_slide(blank); PAGE[0] += 1
    rect(sl, 0, 0, 0.35, SH, fill=HDR)
    tb(sl, 1.0, 3.0, 11, 0.8, title, 32, INK, True)
    tb(sl, 1.0, 3.85, 11, 0.4, sub, 12, GRAY)
    return PAGE[0]

def short_plat(p):
    if 'PC 전용' in p or p.startswith('PC'): return 'PC 전용'
    if '반응형' in p: return '반응형'
    return p[:10]

def index_slide(pages_of, sec_pages):
    sl = prs.slides.add_slide(blank); PAGE[0] += 1
    tb(sl, 0.3, 0.25, 8, 0.4, '화면 목록 (Screen List)', 16, INK, True)
    tb(sl, 0.3, 0.68, 12.7, 0.25, f'PC 시나리오 {sec_pages[0]}p부터, Mobile 시나리오 {sec_pages[1]}p부터. 번호 마커(주황 원)는 각 슬라이드 Description의 구성 요소 번호와 같습니다. 정오·경고·오류는 색이 아니라 문구로 구분합니다.', 8.5, GRAY)
    ids = [s['sid'].upper() for s in geo]; per = 18
    for ti in range(0, len(ids), per):
        chunk = ids[ti:ti + per]; x = 0.3 + (ti // per) * 4.3
        t = new_table(sl, len(chunk) + 1, 5, x, 1.05, 4.1, 0.24 * (len(chunk) + 1))
        for i, w in enumerate([0.8, 1.75, 0.65, 0.45, 0.45]): t.columns[i].width = Inches(w)
        for c, h in enumerate(['Screen ID', '화면명', '구분', 'PC', 'Mobile']): cell_set(t.cell(0, c), h, 7.5, True, WHITE, HDR, PP_ALIGN.CENTER)
        for r, i in enumerate(chunk, 1):
            m = meta.get(i, {})
            cell_set(t.cell(r, 0), i, 7.5, True); cell_set(t.cell(r, 1), m.get('name', ''), 7.5)
            cell_set(t.cell(r, 2), short_plat(m.get('플랫폼', '')), 7, align=PP_ALIGN.CENTER)
            cell_set(t.cell(r, 3), str(pages_of.get((i, 'pc'), '-')), 7.5, align=PP_ALIGN.CENTER)
            cell_set(t.cell(r, 4), str(pages_of.get((i, 'mo'), '-')), 7.5, align=PP_ALIGN.CENTER)
            t.rows[r].height = Inches(0.24)
    return sl

# ---------- 생성 ----------
cover()
prs.slides.add_slide(blank); PAGE[0] += 1   # 화면 목록 자리 (페이지 번호 확정 후 채움)
pages_of = {}
p_pc = divider('PC 화면', 'PC 시나리오 (1000px 기준, PC 전용 화면 포함)')
for scr in geo: render_screen(scr, 'pc', pages_of)
p_mo = divider('Mobile 화면', 'Mobile 시나리오 (375px 기준)')
for scr in geo: render_screen(scr, 'mo', pages_of)
index_slide(pages_of, (p_pc, p_mo))
lst = list(prs.slides._sldIdLst); ids = prs.slides._sldIdLst
ids.remove(lst[-1]); ids.insert(1, lst[-1]); ids.remove(lst[1]); prs.part.drop_rel(lst[1].get(qn('r:id')))
cp = prs.core_properties; now = datetime.datetime.now()
cp.title = f'{PROJ} 화면설계서 {VER}'; cp.author = cp.last_modified_by = ''; cp.description = 'sb-lowfi'; cp.created = cp.modified = now
prs.save(OUT)
print(len(prs.slides), 'slides ->', OUT)
print('marker collisions (no free spot):', len(COLLIDE), COLLIDE[:20])
